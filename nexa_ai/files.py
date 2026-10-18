"""Attachment extraction with size limits and safe ZIP inspection (never extracts ZIPs to disk)."""
import io, os, zipfile
from pathlib import PurePosixPath
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from .config import MAX_ATTACHMENT_MB, MAX_ZIP_TEXT_CHARS

TEXT_EXTS = {
    ".txt", ".md", ".py", ".js", ".jsx", ".ts", ".tsx", ".json", ".html", ".css",
    ".yaml", ".yml", ".toml", ".ini", ".xml", ".sql", ".sh", ".java", ".c", ".cpp",
    ".h", ".go", ".rs", ".php", ".rb", ".swift", ".kt", ".log", ".env.example",
    ".gitignore", ".dockerfile", ".csv"
}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp"}
def validate_size(data: bytes):
    if len(data) > MAX_ATTACHMENT_MB * 1024 * 1024:
        raise ValueError(f"File is too large. Limit is {MAX_ATTACHMENT_MB} MB.")

def extract_attachment(filename: str, data: bytes, mime: str = ""):
    validate_size(data)
    ext = os.path.splitext(filename.lower())[1]
    if ext in IMAGE_EXTS or mime.startswith("image/"):
        return {"kind": "image", "bytes": data, "mime": mime or "image/jpeg",
                "prompt_text": f"[Image attachment: {filename}]"}
    if ext == ".pdf" or mime == "application/pdf":
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        if text.strip():
            return {"kind": "text", "text": f"PDF {filename} contents:\n{text[:MAX_ZIP_TEXT_CHARS]}"}
        return {"kind": "pdf", "bytes": data, "mime": "application/pdf",
                "prompt_text": f"[Scanned or image-based PDF attachment: {filename}]"}
    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return {"kind": "text", "text": f"Document {filename} contents:\n{text[:MAX_ZIP_TEXT_CHARS]}"}
    if ext in {".xlsx", ".xlsm"}:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        chunks = []
        for ws in wb.worksheets:
            chunks.append(f"Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = ["" if v is None else str(v) for v in row]
                if any(vals):
                    chunks.append(" | ".join(vals))
                if sum(map(len, chunks)) > MAX_ZIP_TEXT_CHARS:
                    break
        wb.close()
        return {"kind": "text", "text": f"Spreadsheet {filename} contents:\n" + "\n".join(chunks)[:MAX_ZIP_TEXT_CHARS]}
    if ext == ".zip" or mime == "application/zip":
        return {"kind": "text", "text": inspect_zip(filename, data)}
    if ext in TEXT_EXTS or mime.startswith("text/") or ext in {".doc", ".rtf", ".ppt", ".pptx"}:
        try:
            text = data.decode("utf-8", errors="replace")
        except Exception:
            text = ""
        if text.strip():
            return {"kind": "text", "text": f"File {filename} contents:\n{text[:MAX_ZIP_TEXT_CHARS]}"}
        return {"kind": "text", "text": f"Received {filename}, but automatic text extraction for this binary format is not implemented yet."}
    return {"kind": "text", "text": f"Received attachment {filename} ({mime or 'unknown type'}). Automatic extraction is not supported for this format yet."}

def inspect_zip(filename, data):
    out = [f"ZIP project analysis for {filename}. Files are inspected in memory; no files are extracted to disk."]
    total = 0
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
            out.append(f"Archive contains {len(infos)} entries.")
            for info in infos:
                name = info.filename
                path = PurePosixPath(name)
                if name.startswith(("/", "\\")) or ".." in path.parts or info.is_dir():
                    continue
                if info.file_size > 2_000_000:
                    out.append(f"\n--- {name} (skipped: file too large) ---")
                    continue
                out.append(f"\n--- {name} ({info.file_size} bytes) ---")
                ext = os.path.splitext(name.lower())[1]
                if ext in TEXT_EXTS and not (info.flag_bits & 0x1):
                    try:
                        content = zf.read(info).decode("utf-8", errors="replace")
                        room = MAX_ZIP_TEXT_CHARS - total
                        if room <= 0:
                            out.append("[Text extraction limit reached.]")
                            break
                        content = content[:min(room, 10000)]
                        out.append(content)
                        total += len(content)
                    except Exception:
                        out.append("[Could not read this file.]")
                else:
                    out.append("[Binary or unsupported file; listed but not included.]")
                if total >= MAX_ZIP_TEXT_CHARS:
                    break
    except zipfile.BadZipFile:
        return f"The attachment {filename} is not a valid ZIP archive."
    return "\n".join(out)[:MAX_ZIP_TEXT_CHARS + 12000]
